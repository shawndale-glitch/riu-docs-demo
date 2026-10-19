"""Process inventory workbook -> lint report.

Usage:
  python tooling/inventory_check.py path/to/inventory.xlsx [--sheet NAME] [--report report.md]

Checks the workbook that lists the processes to be redesigned against the `inventory`
rules in governance/workflow-rules.yaml (I1-I8): required columns, process IDs, dates,
statuses, area columns and their key, FigJam links, and agreement with workflows/.

Exit codes: 0 = no errors (warnings allowed), 1 = errors found, 2 = usage error.
"""
from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import workflow_core as core  # noqa: E402

RULES = core.load_yaml(core.REPO / "governance" / "workflow-rules.yaml")["inventory"]
PID = re.compile(r"[A-Z0-9]{2,6}-\d{2,3}(\.\d+)?")
FIGJAM = re.compile(r"https://(www\.)?figma\.com/board/[A-Za-z0-9]+/\S*node-id=\d+[-:]\d+")
KEY_PREFIXES = ("clave:", "key:", "leyenda:")
KEY_SHEETS = {"clave", "key", "leyenda"}


def norm(s) -> str:
    """Case- and accent-insensitive form of a header cell."""
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().casefold()


def check(path: Path, sheet: str | None = None) -> tuple[list[core.Finding], dict]:
    from openpyxl import load_workbook

    findings: list[core.Finding] = []
    err = lambda r, where, l, m: findings.append(core.Finding(r, "error", where, l, m))
    warn = lambda r, where, l, m: findings.append(core.Finding(r, "warning", where, l, m))

    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    aliases = {k: {norm(a) for a in v} for k, v in RULES["required_columns"].items()}

    # Header row: the first row (within 10) naming at least two required columns.
    header_row, cols = None, {}
    for r in range(1, min(ws.max_row, 10) + 1):
        found = {}
        for c in range(1, ws.max_column + 1):
            h = norm(ws.cell(r, c).value)
            for key, names in aliases.items():
                if h in names and key not in found:
                    found[key] = c
        if len(found) >= 2:
            header_row, cols = r, found
            break
    if header_row is None:
        err("I1", None, "", "No header row found: none of the first 10 rows names two required columns")
        return findings, {"sheet": ws.title, "rows": 0}
    for key in RULES["required_columns"]:
        if key not in cols:
            err("I1", f"row {header_row}", key, f"Missing column '{RULES['required_columns'][key][0]}' (accepted names: {', '.join(RULES['required_columns'][key])})")

    headers = {c: ws.cell(header_row, c).value for c in range(1, ws.max_column + 1) if ws.cell(header_row, c).value not in (None, "")}
    ignored = {norm(x) for x in RULES.get("ignored_columns", [])}
    area_cols = [c for c in headers if c not in cols.values() and norm(headers[c]) not in ignored]

    # Data rows: non-empty rows below the header, excluding key rows.
    rows, key_found = [], ws.title.casefold() in KEY_SHEETS or any(s.title.casefold() in KEY_SHEETS for s in wb.worksheets)
    for r in range(header_row + 1, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if all(v in (None, "") for v in values):
            continue
        first = next((v for v in values if v not in (None, "")), "")
        if str(first).strip().casefold().startswith(KEY_PREFIXES):
            key_found = True
            continue
        rows.append(r)

    cell = lambda r, key: ws.cell(r, cols[key]) if key in cols else None
    seen: dict[str, int] = {}
    inventory: dict[str, str] = {}
    for r in rows:
        where = f"row {r}"
        name = str(cell(r, "process").value or "") if "process" in cols else ""
        c = cell(r, "process_id")
        if c is not None:
            v = c.value
            if v in (None, ""):
                err("I2", where, name, "Process ID is empty")
            else:
                pid = str(v).strip()
                if c.data_type == "n":
                    err("I2", where, pid, "Process ID is stored as a number; store it as text (e.g. 'RP-01')")
                if not PID.fullmatch(pid):
                    err("I2", where, pid, "Process ID must look like AREA-NN or AREA-NNN, sub-process AREA-NN.n (e.g. RP-01, BB-02)")
                elif pid.split("-")[0] not in core.areas():
                    err("I2", where, pid, f"Unknown area prefix '{pid.split('-')[0]}'; add it to governance/areas.yaml")
                if pid in seen:
                    err("I2", where, pid, f"Duplicate process ID (also on row {seen[pid]})")
                seen.setdefault(pid, r)
        status = None
        c = cell(r, "status")
        if c is not None:
            raw = str(c.value or "").strip()
            status = core.STATUS_MAP.get(raw)
            if not status:
                err("I4", where, raw, f"Status must be one of {list(core.STATUS_MAP)}")
            elif c is not None and cell(r, "process_id") is not None and cell(r, "process_id").value:
                inventory[str(cell(r, "process_id").value).strip()] = status
        c = cell(r, "planned_date")
        if c is not None and c.value not in (None, ""):
            v = c.value
            if not isinstance(v, (date, datetime)) and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(v).strip()):
                err("I3", where, str(v), "Planned date must be a date cell or YYYY-MM-DD text")
        c = cell(r, "figjam_link")
        if c is not None:
            link = str((c.hyperlink.target if c.hyperlink else None) or c.value or "").strip()
            if not FIGJAM.match(link):
                msg = "FigJam link missing or not a link to a process section (figma.com/board/...?node-id=...)"
                (err if status in ("in-review", "validated") else warn)("I7", where, link[:60] or name, msg)

    for c in area_cols:
        if not any(ws.cell(r, c).value not in (None, "") for r in rows):
            warn("I6", f"column {ws.cell(header_row, c).column_letter}", str(headers[c]), "Column has a header but no values; fill it or remove it")
    used_area_cols = [c for c in area_cols if any(ws.cell(r, c).value not in (None, "") for r in rows)]
    if used_area_cols and not key_found:
        names = ", ".join(str(headers[c]) for c in used_area_cols)
        warn("I5", None, names, "Area columns have no key: add a row starting 'Clave:' (or a 'Clave' sheet) saying what a mark means, e.g. 'Clave: X = participa'")

    for wf_path in sorted((core.REPO / "workflows").glob("*/*/workflow.yaml")):
        p = core.load_yaml(wf_path)["process"]
        if p["id"] not in seen:
            warn("I8", None, p["id"], f"Process in {wf_path.relative_to(core.REPO).as_posix()} is missing from the workbook")
        elif inventory.get(p["id"]) and inventory[p["id"]] != p["status"]:
            warn("I8", f"row {seen[p['id']]}", p["id"], f"Workbook status '{inventory[p['id']]}' differs from workflow.yaml status '{p['status']}'")

    return findings, {"sheet": ws.title, "header_row": header_row, "rows": len(rows), "area_columns": [str(headers[c]) for c in area_cols]}


def report(path: Path, findings: list[core.Finding], info: dict) -> str:
    errors = [f for f in findings if f.severity == "error"]
    warnings = [f for f in findings if f.severity == "warning"]
    L = ["# Process inventory check", "",
         f"- Workbook `{path.name}`, sheet `{info['sheet']}`, header on row {info.get('header_row', '?')}, {info['rows']} process rows",
         f"- Area columns: {', '.join(info.get('area_columns') or []) or 'none'}", "",
         f"## Findings: {len(errors)} errors, {len(warnings)} warnings", ""]
    if not findings:
        L.append("None.")
    for f in errors + warnings:
        where = f" ({f.node})" if f.node else ""
        L.append(f"- {'❌' if f.severity == 'error' else '⚠️'} **{f.rule}**{where} «{f.label}» {f.message}")
    L += ["", f"Rules: `governance/workflow-rules.yaml` → `inventory` (rules_version {core.load_yaml(core.REPO / 'governance' / 'workflow-rules.yaml')['rules_version']})."]
    return "\n".join(L) + "\n"


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("workbook")
    ap.add_argument("--sheet", default=None)
    ap.add_argument("--report", default=None)
    a = ap.parse_args()
    path = Path(a.workbook)
    if not path.exists():
        print(f"Not found: {path}", file=sys.stderr)
        return 2
    findings, info = check(path, a.sheet)
    text = report(path, findings, info)
    if a.report:
        Path(a.report).write_text(text, encoding="utf-8")
    print(text)
    return 1 if any(f.severity == "error" for f in findings) else 0


if __name__ == "__main__":
    sys.exit(main())
